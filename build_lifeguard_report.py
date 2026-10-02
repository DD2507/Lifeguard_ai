from pathlib import Path
from datetime import date
from PIL import Image
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(r"C:\Users\Dhrupad\Desktop\Major project mobileapp")
ASSETS = ROOT / "report_assets"
OUT = ROOT / "LifeGuard_AI_Application_Report.docx"


def crop(src, dst, top_frac=0.0, bottom_frac=1.0, ratio=1.48):
    im = Image.open(src)
    y1, y2 = int(im.height * top_frac), int(im.height * bottom_frac)
    target_h = int(im.width / ratio)
    y2 = min(y2, y1 + target_h)
    im.crop((0, y1, im.width, y2)).save(dst, quality=92)


crop(ASSETS / "dashboard.jpg", ASSETS / "dashboard_top.jpg", 0, 0.72)
crop(ASSETS / "patients.jpg", ASSETS / "patients_top.jpg", 0, 0.70)
crop(ASSETS / "alerts.jpg", ASSETS / "alerts_top.jpg", 0, 0.85)
crop(ASSETS / "rooms.jpg", ASSETS / "rooms_top.jpg", 0, 0.75)
crop(ASSETS / "beds.jpg", ASSETS / "beds_top.jpg", 0, 0.55)
crop(ASSETS / "patient_detail.jpg", ASSETS / "patient_detail_top.jpg", 0, 0.34)
crop(ASSETS / "patient_detail.jpg", ASSETS / "patient_detail_mid.jpg", 0.28, 0.64)
crop(ASSETS / "patient_detail.jpg", ASSETS / "patient_detail_rx.jpg", 0.60, 0.93)

doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.7)
sec.bottom_margin = Inches(0.65)
sec.left_margin = Inches(0.78)
sec.right_margin = Inches(0.78)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(10.5)
styles["Normal"].paragraph_format.space_after = Pt(6)
styles["Normal"].paragraph_format.line_spacing = 1.12
for style_name, size in [("Title", 30), ("Heading 1", 19), ("Heading 2", 14), ("Heading 3", 11.5)]:
    s = styles[style_name]
    s.font.name = "Aptos Display"
    s.font.size = Pt(size)
    s.font.color.rgb = RGBColor(0, 0, 0)
    s.font.bold = True
    s.paragraph_format.space_before = Pt(10)
    s.paragraph_format.space_after = Pt(6)

# Remove inherited paragraph borders from the built-in Title style.
title_ppr = styles["Title"]._element.get_or_add_pPr()
for border in title_ppr.findall(qn("w:pBdr")):
    title_ppr.remove(border)


def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def set_cell_margins(cell, top=110, start=120, bottom=110, end=120):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tcMar.append(node)
        node.set(qn("w:w"), str(v)); node.set(qn("w:type"), "dxa")


def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = "Table Grid"
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = h
        set_cell_shading(c, "16324F")
        for r in c.paragraphs[0].runs:
            r.font.color.rgb = RGBColor(255,255,255); r.font.bold = True; r.font.size = Pt(9)
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(c)
    for ridx, row in enumerate(rows):
        cells = t.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)
            if ridx % 2: set_cell_shading(cells[i], "EFF5F8")
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cells[i])
            for p in cells[i].paragraphs:
                p.paragraph_format.space_after = Pt(0)
                for r in p.runs: r.font.size = Pt(8.7)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths): row.cells[i].width = Inches(w)
    doc.add_paragraph()
    return t


def bullet(text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.add_run(text)
    return p


def figure(path, caption, width=6.75):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(path), width=Inches(width))
    c = doc.add_paragraph(caption)
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c.paragraph_format.space_after = Pt(10)
    for r in c.runs:
        r.italic = True; r.font.size = Pt(9); r.font.color.rgb = RGBColor(70,70,70)


def mobile_pair(left_path, left_caption, right_path, right_caption):
    t = doc.add_table(rows=2, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    tblPr = t._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}"); el.set(qn("w:val"), "nil"); borders.append(el)
    tblPr.append(borders)
    for idx, path in enumerate((left_path, right_path)):
        p = t.rows[0].cells[idx].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(path), width=Inches(2.65))
        set_cell_margins(t.rows[0].cells[idx], top=30, start=70, bottom=30, end=70)
    for idx, caption in enumerate((left_caption, right_caption)):
        p = t.rows[1].cells[idx].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(caption); r.italic = True; r.font.size = Pt(8.5); r.font.color.rgb = RGBColor(70,70,70)
        set_cell_margins(t.rows[1].cells[idx], top=25, start=70, bottom=60, end=70)


def page_break(): doc.add_page_break()


# Cover
p = doc.add_paragraph(style="Title")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("LifeGuard AI Application Report")
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Intelligent patient monitoring, environmental sensing, clinical alerts, and digital twin simulation")
r.font.size = Pt(14); r.font.color.rgb = RGBColor(45,72,88)
doc.add_paragraph("\n")
figure(ASSETS / "dashboard_top.jpg", "Working LifeGuard AI dashboard captured from the locally running application", 6.5)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("Prepared from the complete project repository and a verified local runtime\n").bold = True
p.add_run("2 October 2026")
page_break()

doc.add_heading("Executive Summary", level=1)
doc.add_paragraph(
    "LifeGuard AI is a multi-platform healthcare monitoring prototype designed for continuous observation of patients and their rooms. "
    "It combines ESP32-connected sensors, MQTT telemetry, a Node.js and MongoDB backend, a Python machine-learning inference service, a React clinical dashboard, and a Flutter mobile client. The system converts incoming physiological and environmental readings into validated patient state, personalized baselines, explainable risk assessments, alerts, room-control recommendations, medicine reminders, and a patient digital twin."
)
doc.add_paragraph(
    "The strongest aspect of the project is its end-to-end integration. A sensor reading is not merely displayed: it is checked for validity and freshness, compared with a patient-specific baseline, scored by rules and optionally by a trained random-forest model, persisted, broadcast in real time, and surfaced through web and mobile interfaces. The application also addresses operational needs through bed management, room monitoring, prescriptions, reminders, and alert history."
)
table(["Area", "Verified project capability"], [
    ("Patient monitoring", "Heart rate, SpO2, body temperature, sensor contact, freshness, and baseline status"),
    ("Room monitoring", "Temperature, humidity, raw MQ135 air-quality signal, room sensor status, and control actions"),
    ("Clinical intelligence", "Rule-based risk, trained AI inference, baseline deviation, explanations, trends, and what-if simulation"),
    ("Clinical workflow", "Patient list, alerts, prescriptions, medicine reminders, bed assignment, release, and maintenance states"),
    ("Real-time delivery", "MQTT ingestion plus Socket.IO updates to dashboards and notification workflows"),
    ("Quality evidence", "React production build passed; 7 Flutter tests and 42 backend tests passed"),
], [1.55, 5.0])

doc.add_heading("Report Scope", level=1)
doc.add_paragraph("This report describes the application as implemented in the current repository, including the active backend, web dashboard, Flutter client, machine-learning service, persistence model, MQTT integration, testing evidence, runtime observations, limitations, and recommended future work. Screenshots were captured from the locally running React application connected to the live backend and database.")

doc.add_heading("Contents", level=1)
content_items = ["1 Project Purpose and Users", "2 System Architecture", "3 Major Features", "4 User Interface Walkthrough", "5 Backend and API Design", "6 AI and Risk Assessment", "7 Data Model and Real-Time Flow", "8 Technology Stack", "9 Testing and Runtime Verification", "10 Security and Safety Considerations", "11 Limitations and Future Enhancements", "12 Conclusion"]
contents_table = doc.add_table(rows=6, cols=2)
contents_table.alignment = WD_TABLE_ALIGNMENT.LEFT
tblPr = contents_table._tbl.tblPr
borders = OxmlElement("w:tblBorders")
for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
    el = OxmlElement(f"w:{edge}"); el.set(qn("w:val"), "nil"); borders.append(el)
tblPr.append(borders)
for i in range(6):
    for j in range(2):
        cell = contents_table.rows[i].cells[j]
        cell.text = content_items[i + j * 6]
        set_cell_margins(cell, top=30, start=0, bottom=30, end=80)
        for r in cell.paragraphs[0].runs: r.font.size = Pt(8.3)
page_break()

doc.add_heading("1 Project Purpose and Users", level=1)
doc.add_heading("1.1 Problem addressed", level=2)
doc.add_paragraph("Hospitals need a concise view of changing patient condition, sensor reliability, room environment, medication timing, and bed availability. These signals normally come from separate devices and workflows. LifeGuard AI brings them into one operational interface and emphasizes early detection, personalization, and explainability rather than presenting raw sensor values alone.")
doc.add_heading("1.2 Intended users", level=2)
table(["User", "Primary tasks"], [
    ("Nurses and ward staff", "Monitor live vitals, acknowledge alerts, complete medicine doses, inspect rooms, and track bed status"),
    ("Doctors", "Review patient detail, risk factors, trends, digital-twin output, and create or manage prescriptions"),
    ("Hospital operations staff", "Observe capacity, assign or release beds, and mark beds for maintenance"),
    ("System administrators and developers", "Run services, verify sensor feeds, maintain models, inspect APIs, and extend the platform"),
], [1.8, 4.75])
doc.add_heading("1.3 Core objectives", level=2)
for x in [
    "Collect physiological readings from patient hardware and environmental readings from room hardware.",
    "Reject invalid, disconnected, stale, retained, or unsafe calibration samples instead of inventing measurements.",
    "Build patient-specific baselines from valid normal MQTT readings and calculate signed deviations.",
    "Generate understandable risk levels, reasons, and response recommendations.",
    "Deliver alerts and scheduled medicine reminders even when no dashboard is currently connected.",
    "Provide consistent web and mobile clinical views backed by the same REST and real-time services.",
]: bullet(x)

doc.add_heading("2 System Architecture", level=1)
doc.add_paragraph("LifeGuard AI follows a layered, event-driven architecture. The physical layer publishes sensor data. The backend validates and enriches it. Persistent models store the clinical and operational state. Intelligence services assess risk. REST and Socket.IO expose this state to the React and Flutter clients.")
table(["Layer", "Components", "Responsibility"], [
    ("Sensing", "Patient ESP32, room ESP32, pulse oximeter, temperature and room sensors", "Measure patient and room conditions and publish telemetry"),
    ("Messaging", "Mosquitto MQTT broker", "Carry device events to backend topics with low latency"),
    ("Application backend", "Node.js, Express, Socket.IO", "Validate data, apply business rules, persist records, schedule reminders, and expose APIs"),
    ("Intelligence", "Python FastAPI, scikit-learn random forest, SHAP artifact", "Perform in-domain risk prediction and provide model metadata/explanations"),
    ("Persistence", "MongoDB with Mongoose models", "Store patients, vitals, rooms, alerts, prescriptions, and beds"),
    ("Clients", "React dashboard and Flutter application", "Present monitoring, alerts, digital twin, treatments, rooms, and bed workflows"),
], [1.15, 2.1, 3.35])
doc.add_heading("2.1 End-to-end data flow", level=2)
for i, x in enumerate([
    "ESP32 devices publish patient or room payloads through MQTT.",
    "The MQTT service maps topics to patients or rooms and forwards payloads to validation services.",
    "Patient and room services reject malformed, conflicting, stale, retained, disconnected, or out-of-range input.",
    "Valid patient data updates the live record and calibration window; valid room data updates environmental context.",
    "The risk engine calculates an explainable rule result; AI inference is requested only when all nine model inputs are finite and inside the training domain.",
    "Vitals, patient state, risk evidence, room state, and alerts are persisted in MongoDB.",
    "Socket.IO broadcasts patient, room, alert, and prescription events to connected clients.",
    "React and Flutter refresh or react to real-time events and display the latest clinical state.",
], 1):
    p=doc.add_paragraph(style="List Number"); p.add_run(x)
page_break()

doc.add_heading("3 Major Features", level=1)
features = [
    ("Live patient telemetry", "Displays heart rate, oxygen saturation, body temperature, room assignment, sensor state, last valid timestamp, risk state, and baseline progress. The UI explicitly distinguishes valid values from unavailable data."),
    ("Sensor integrity and freshness", "Tracks finger contact, sensor disconnection, invalid values, retained MQTT messages, stale timestamps, and last valid readings. Invalid packets preserve the last trustworthy measurement without entering baseline calibration."),
    ("Personalized baseline calibration", "Collects valid normal readings, uses a robust median for the main patient baseline, stores sample counts and status, freezes an established baseline, and computes signed heart-rate, SpO2, and temperature deviations."),
    ("Explainable hybrid risk scoring", "Combines rule-based screening with optional random-forest inference. The system records risk, numeric score, reasons, model source, model version, confidence/probabilities, and input-domain issues."),
    ("Patient digital twin", "Builds a structured summary of current state, personal baseline, deviation, and trend. A what-if simulator applies reduced oxygen, fever-like, elevated heart-rate, combined deterioration, or poor-room scenarios without changing live data."),
    ("Alerts and history", "Creates clinical alerts for qualifying episodes, prevents duplicate high-BPM alerts during the same episode, resolves alerts when conditions normalize, retains history, and emits real-time changes."),
    ("Prescription and medicine reminders", "Creates, updates, completes, reactivates, and deletes treatments. A scheduler resolves India-time schedules, persists due alerts, emits reminders once, and retries safely when Socket.IO is unavailable."),
    ("Room monitoring and control", "Displays temperature, humidity, air-quality sensor values, readiness and status. Backend control routes publish fan or buzzer actions through MQTT and return a recommended room response."),
    ("Bed management", "Shows total, available, occupied, and maintenance capacity; supports patient assignment, release, maintenance, restoration, creation, and deletion through dedicated APIs."),
    ("Real-time multi-client experience", "The web dashboard polls frequently and listens for socket events; the Flutter notification service handles patient, room, alert, and medicine events, including local notifications and recovery of active medicine alerts."),
]
for title, body in features:
    doc.add_heading(title, level=2); doc.add_paragraph(body)
doc.add_heading("4 User Interface Walkthrough", level=1)
doc.add_paragraph("The React dashboard uses a clinical operations layout with a persistent navigation bar, system-status indicator, summary metrics, searchable patient lists, and detailed operational pages. The following figures were captured from the running application on 2 October 2026.")
figure(ASSETS / "dashboard_top.jpg", "Figure 1 Dashboard overview with census, risk summary, active-alert status, and patient monitoring", 6.7)
figure(ASSETS / "patients_top.jpg", "Figure 2 Patient directory with search, risk filters, sensor state, baseline progress, and live vitals", 6.7)
page_break()
figure(ASSETS / "alerts_top.jpg", "Figure 3 Alert workspace for active clinical and medicine alerts with alert-history context", 6.7)
figure(ASSETS / "rooms_top.jpg", "Figure 4 Room monitoring view combining environmental readings and assigned patients", 6.7)
page_break()
figure(ASSETS / "beds_top.jpg", "Figure 5 Bed availability and occupancy management with capacity metrics and status filters", 6.7)
doc.add_heading("4.1 Patient detail and digital twin", level=2)
doc.add_paragraph("The patient-detail view consolidates live vital signs, room environment, sensor health, baseline collection, risk assessment, explainable factors, room-action advice, digital-twin state, simulations, and prescriptions. This is the most complete clinical view in the application.")
figure(ASSETS / "patient_detail_top.jpg", "Figure 6 Patient vital signs, room environment, and baseline collection", 6.7)
page_break()
figure(ASSETS / "patient_detail_mid.jpg", "Figure 7 Digital-twin state, deviations, trends, simulation, and explainable risk", 6.7)
figure(ASSETS / "patient_detail_rx.jpg", "Figure 8 Prescription management and treatment status on the patient record", 6.7)

page_break()
doc.add_heading("4.2 Flutter mobile application", level=2)
doc.add_paragraph("The Flutter client was repaired and run in a 390 × 844 mobile viewport against the live Node.js API. It provides a compact clinical workflow with the same patient, alert, and room information used by the web dashboard. The bottom navigation keeps the four principal monitoring areas within one tap.")
mobile_pair(
    ASSETS / "mobile" / "flutter_dashboard.jpg",
    "Figure 9 Flutter dashboard with live census, alert state, and patient vitals",
    ASSETS / "mobile" / "flutter_patients.jpg",
    "Figure 10 Flutter patient list with room, risk, heart rate, SpO2, and temperature",
)
page_break()
mobile_pair(
    ASSETS / "mobile" / "flutter_alerts.jpg",
    "Figure 11 Flutter alerts screen showing active monitoring status",
    ASSETS / "mobile" / "flutter_rooms.jpg",
    "Figure 12 Flutter smart-room view with temperature, humidity, air quality, and device state",
)

doc.add_heading("5 Backend and API Design", level=1)
doc.add_paragraph("The active backend is the `backend` Node.js application on port 5000. It loads environment configuration, connects to MongoDB, starts the MQTT subscriber, registers Socket.IO, mounts domain routers, and starts the prescription reminder scheduler. The older `server` folder appears to be a previous or simplified implementation and is not the startup target documented by the project.")
table(["API group", "Representative operations"], [
    ("Patients", "Create/list/detail; establish or reset baseline; retrieve digital twin; run what-if simulation"),
    ("Vitals", "Submit patient readings and retrieve vital history"),
    ("Alerts", "List active alerts, history, patient alerts, and resolve/complete alert workflows"),
    ("Rooms", "List/detail rooms, receive environment data, and execute control actions"),
    ("Prescriptions", "Create, list by patient, retrieve, update, change status, and delete"),
    ("Beds", "List, filter available, create, update, assign, release, maintenance, available, and delete"),
], [1.55, 5.0])
doc.add_heading("5.1 Service responsibilities", level=2)
table(["Service", "Role"], [
    ("patientVitalService", "Serializes updates per patient; validates source and sensor state; manages baseline, risk, alerts, persistence, and broadcasts"),
    ("baselineService", "Validates readings, calculates median or mean baselines, checks persistence integrity, and calculates signed deviations"),
    ("riskEngine", "Calculates rule score, risk level, reasons, recommendations, and room actions"),
    ("aiClient", "Builds the exact nine-feature vector, validates the model domain, calls FastAPI, and handles unavailable inference"),
    ("digitalTwinService", "Computes history-based baseline/trends and builds current and simulated patient summaries"),
    ("roomSensorService", "Parses room payload aliases, validates fields and room identity, preserves last valid state, and emits updates"),
    ("prescriptionReminderScheduler", "Claims due reminders, persists medicine alerts, emits once, retries failure, and handles India-time schedules"),
    ("notificationService", "Registers Socket.IO rooms and emits typed patient, room, alert, and prescription events"),
], [2.05, 4.5])

doc.add_heading("6 AI and Risk Assessment", level=1)
doc.add_heading("6.1 Model", level=2)
doc.add_paragraph("The Python service exposes prediction through FastAPI on port 8000. It loads a trained `RandomForestClassifier`, model metadata, and a SHAP explainer artifact. The model expects nine features: current heart rate, SpO2, body temperature, room temperature, humidity, air quality, and signed deviations for heart rate, SpO2, and temperature.")
doc.add_heading("6.2 Guardrails", level=2)
for x in [
    "Inference is skipped when any required value is missing, non-finite, or outside the recorded training range.",
    "The rule engine remains available when the AI service is unreachable or the input is out of domain.",
    "Sensor-invalid states withhold the risk score instead of presenting a misleading result.",
    "The what-if simulator is clearly labeled as an estimate and not a diagnosis.",
    "The project includes a model card and dataset requirements describing expected validation and limitations.",
]: bullet(x)
doc.add_heading("6.3 Baseline and risk interaction", level=2)
doc.add_paragraph("Calibration accepts only valid, fresh, normal-range MQTT readings. This prevents abnormal episodes, HTTP test values, retained packets, or disconnected sensor values from being absorbed into the patient's normal baseline. Once enough samples are collected, the baseline becomes stable and deviations help identify meaningful changes for that individual.")

doc.add_heading("7 Data Model and Real-Time Flow", level=1)
table(["Entity", "Important stored information"], [
    ("Patient", "Identity, room, current vitals, sensor status, baseline samples/status, risk result, AI metadata, and timestamps"),
    ("Vital", "Patient reading history, source, sensor validity, risk fields, room context, and received time"),
    ("Room", "Environmental values, sensor status, last valid/read times, patient relationship, fan and buzzer state"),
    ("Alert", "Patient, type, severity, message, status, clinical/medicine context, due time, and completion/resolution data"),
    ("Prescription", "Patient, treatment, dose, schedule, status, reminder claim/delivery state, and completion data"),
    ("Bed", "Bed identifier, room, ward, status, assigned patient, and maintenance/occupancy metadata"),
], [1.4, 5.15])
doc.add_heading("7.1 Real-time event categories", level=2)
doc.add_paragraph("Socket.IO events cover prescription creation, medicine due, patient vital updates, room updates, room-sensor status, alert creation/update/completion, and alert resolution. Clients can join global monitoring rooms or patient-specific rooms. This allows both a ward dashboard and an individual patient view to remain synchronized.")
doc.add_heading("7.2 MQTT integration", level=2)
doc.add_paragraph("The MQTT service connects to the broker on port 1883, subscribes to patient and room topics, maps payloads into domain services, and publishes room-control commands. The project notes indicate separate room and patient ESP32 devices, with room 102 and patient P003 used as the primary demonstration path.")

doc.add_heading("8 Technology Stack", level=1)
table(["Category", "Technology"], [
    ("Mobile application", "Flutter and Dart; HTTP, Socket.IO client, local notifications, and timezone support"),
    ("Web application", "React 19, Vite, Socket.IO client, responsive CSS"),
    ("Backend", "Node.js, Express, Mongoose, Socket.IO, MQTT, Axios, dotenv"),
    ("AI service", "Python, FastAPI, Uvicorn, scikit-learn, joblib, NumPy/Pandas, SHAP artifact"),
    ("Database", "MongoDB"),
    ("IoT messaging", "Mosquitto MQTT"),
    ("Hardware concept", "ESP32 patient and room nodes with physiological and environmental sensors"),
    ("Testing", "Node built-in test runner, Flutter test, Vite production build"),
], [1.75, 4.8])

doc.add_heading("9 Testing and Runtime Verification", level=1)
table(["Verification", "Result", "Evidence"], [
    ("Backend service tests", "PASS", "42 tests passed covering digital twin, AI-domain checks, baselines, sensor validation, alerts, reminders, and room parsing"),
    ("Flutter tests", "PASS", "7 tests passed covering baseline collection, medicine-due UI, and login shell"),
    ("React production build", "PASS", "57 modules transformed; optimized build created successfully"),
    ("Live backend", "PASS", "Node server connected to MongoDB on port 5000"),
    ("AI service", "RUNNING WITH WARNINGS", "FastAPI started on port 8000; model loaded with a scikit-learn version mismatch warning"),
    ("MQTT broker", "PASS", "Mosquitto was listening locally on port 1883"),
    ("React runtime", "PASS", "Dashboard logged in and all principal pages rendered against live API data"),
    ("Flutter Windows launch", "BLOCKED BY HOST SETTING", "Plugin build requires Windows Developer Mode symlink support"),
    ("Flutter web launch", "PASS AFTER REPAIR", "Added the Flutter SDK web-plugin dependency and used a platform-aware localhost API URL"),
], [1.85, 1.35, 3.35])
doc.add_heading("9.1 Observed runtime state", level=2)
doc.add_paragraph("The running database contained four active patients and ten beds. Patient P003 had live patient and room values but reported invalid pulse-oximeter finger contact, so the interface correctly withheld the risk score. The bed view reported seven available beds, one occupied bed, and two maintenance beds. These values demonstrate live rendering and also show the project's explicit handling of incomplete or invalid clinical input.")

page_break()
doc.add_heading("10 Security and Safety Considerations", level=1)
doc.add_paragraph("This is a prototype and should not be represented as a certified clinical system. Several design choices improve safety, while production deployment requires additional controls.")
table(["Current strength", "Production requirement"], [
    ("Invalid and stale sensor states are surfaced", "Define clinical validation protocols, device certification, and alarm-response procedures"),
    ("AI can fall back to explainable rules", "Lock model/library versions, validate prospectively, monitor drift, and document approval boundaries"),
    ("Simulation is labeled non-diagnostic", "Add formal clinical disclaimers, human review gates, and audit trails"),
    ("Alerts and reminder events are persisted", "Add role-based access, immutable audit logging, escalation policies, and delivery monitoring"),
    ("Data is centralized through APIs", "Use authenticated TLS, secrets management, authorization, encryption at rest, and least-privilege network access"),
    ("Demo login is clearly visible", "Remove hard-coded demo credentials and implement secure identity, session expiry, and account lifecycle management"),
], [3.15, 3.05])

page_break()
doc.add_heading("11 Limitations and Future Enhancements", level=1)
doc.add_heading("11.1 Current limitations", level=2)
for x in [
    "The visible login is a demonstration mechanism and is not a production authentication system.",
    "The trained model artifact was created with scikit-learn 1.7.2 but the active environment uses 1.9.0, producing compatibility warnings.",
    "Flutter desktop cannot build on the current host until Developer Mode or equivalent symlink support is enabled.",
    "The Flutter web target now runs, but the native Windows target still depends on host-level symlink support.",
    "The project contains both `backend` and an older `server` directory, which can confuse deployment and maintenance.",
    "Clinical performance evidence, prospective validation, authorization, encrypted deployment, and regulatory controls are not yet complete.",
]: bullet(x)
doc.add_heading("11.2 Recommended next steps", level=2)
for x in [
    "Pin Python and scikit-learn versions to the model-training environment, retrain/export if necessary, and add a reproducible model pipeline.",
    "Repair Flutter platform configuration, verify Android emulator/device execution, and capture mobile-native regression screenshots.",
    "Consolidate the legacy server folder or document it explicitly as archived code.",
    "Implement authenticated roles for nurses, doctors, administrators, and device publishers with auditable actions.",
    "Add API integration tests, end-to-end browser tests, MQTT contract tests, and CI gates for all three application layers.",
    "Add trend charts, alert escalation/acknowledgement, device fleet health, model drift monitoring, and exportable clinical summaries.",
    "Conduct usability and safety testing with representative clinical users before any real-world use.",
]: bullet(x)

doc.add_heading("12 Conclusion", level=1)
doc.add_paragraph("LifeGuard AI is a substantial full-stack and IoT healthcare prototype. It successfully unifies patient telemetry, room conditions, personalized baseline calibration, explainable risk assessment, real-time alerts, prescription reminders, bed operations, and patient digital-twin simulation. The repository demonstrates meaningful engineering depth through data validation, fallback behavior, persistent scheduling, real-time messaging, and automated tests.")
doc.add_paragraph("The application is well suited for an academic major-project demonstration because it shows integration across hardware, messaging, backend services, artificial intelligence, databases, web, and mobile development. Its next phase should focus on platform reproducibility, security, clinical validation, and deployment discipline so that the prototype can evolve into a safer and more maintainable system.")

# Footer with page number field
for section in doc.sections:
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("LifeGuard AI Application Report  |  ")
    r.font.size = Pt(8); r.font.color.rgb = RGBColor(90,90,90)
    fldChar1 = OxmlElement('w:fldChar'); fldChar1.set(qn('w:fldCharType'), 'begin')
    instrText = OxmlElement('w:instrText'); instrText.set(qn('xml:space'), 'preserve'); instrText.text = ' PAGE '
    fldChar2 = OxmlElement('w:fldChar'); fldChar2.set(qn('w:fldCharType'), 'end')
    r._r.append(fldChar1); r._r.append(instrText); r._r.append(fldChar2)

doc.core_properties.title = "LifeGuard AI Application Report"
doc.core_properties.subject = "Detailed project report with verified application screenshots"
doc.core_properties.author = "Project Team"
doc.save(OUT)
print(OUT)
